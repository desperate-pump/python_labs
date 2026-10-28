"""Простая демонстрация трёх заданий ЛР5."""

import sys
from pathlib import Path

from openpyxl import load_workbook

from src.lab05.json_csv import json_to_csv, csv_to_json
from src.lab05.csv_xlsx import csv_to_xlsx


def main() -> int:
    if len(sys.argv) != 2 or sys.argv[1] not in ("json-to-csv", "csv-to-json", "csv-to-xlsx"):
        print("Выберите: json-to-csv, csv-to-json или csv-to-xlsx", file=sys.stderr)
        return 1

    try:
        if sys.argv[1] == "json-to-csv":
            json_to_csv("data/samples/people.json", "data/out/people_from_json.csv")
            print("JSON → CSV")
            print(Path("data/out/people_from_json.csv").read_text(encoding="utf-8"))
        elif sys.argv[1] == "csv-to-json":
            csv_to_json("data/samples/people.csv", "data/out/people_from_csv.json")
            print("CSV → JSON")
            print(Path("data/out/people_from_csv.json").read_text(encoding="utf-8"))
        else:
            csv_to_xlsx("data/samples/people.csv", "data/out/people.xlsx")
            workbook = load_workbook("data/out/people.xlsx")
            sheet = workbook.active
            print("CSV → XLSX")
            print("Лист:", sheet.title)
            print("Строк:", sheet.max_row, "Столбцов:", sheet.max_column)
            for column in sheet.columns:
                letter = column[0].column_letter
                print(f"Ширина {letter}: {sheet.column_dimensions[letter].width:g}")
            for row in sheet.values:
                print(" | ".join(str(value or "") for value in row))
            workbook.close()
    except (OSError, ValueError) as error:
        print("Ошибка:", error, file=sys.stderr)
        return 1

    print("Готово")
    return 0


if __name__ == "__main__":
    sys.exit(main())
