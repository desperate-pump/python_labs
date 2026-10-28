"""Проверки заданий A и B лабораторной №5."""

import csv
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from src.lab05.json_csv import csv_to_json, json_to_csv
from src.lab05.csv_xlsx import csv_to_xlsx


class TestLab05(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.json_path = self.folder / "people.json"
        self.csv_path = self.folder / "people.csv"
        self.xlsx_path = self.folder / "people.xlsx"

    def test_json_to_csv(self):
        data = [{"name": "Анна", "age": 18}, {"name": "Тимур", "city": "Избербаш"}]
        self.json_path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        json_to_csv(str(self.json_path), str(self.csv_path))
        with self.csv_path.open(encoding="utf-8", newline="") as file:
            self.assertEqual(list(csv.reader(file)), [
                ["age", "city", "name"], ["18", "", "Анна"], ["", "Избербаш", "Тимур"]
            ])

    def test_csv_to_json(self):
        self.csv_path.write_text('name,age\n"Мария, Маша",001\n', encoding="utf-8")
        csv_to_json(str(self.csv_path), str(self.json_path))
        self.assertEqual(json.loads(self.json_path.read_text(encoding="utf-8")),
                         [{"name": "Мария, Маша", "age": "001"}])

    def test_empty_and_wrong_json(self):
        for value in ("", "[]", "{}", "[1]", "[{}]", '[{"x": null}]', "{"):
            with self.subTest(value=value):
                self.json_path.write_text(value, encoding="utf-8")
                with self.assertRaises(ValueError):
                    json_to_csv(str(self.json_path), str(self.csv_path))

    def test_invalid_csv(self):
        for value in ("", "\n", "a,\n", "a,a\n", "a,b\n1\n", "a,b\n1,2,3\n", 'a,b\n"bad,3\n'):
            with self.subTest(value=value):
                self.csv_path.write_text(value, encoding="utf-8")
                with self.assertRaises(ValueError):
                    csv_to_json(str(self.csv_path), str(self.json_path))
                with self.assertRaises(ValueError):
                    csv_to_xlsx(str(self.csv_path), str(self.xlsx_path))

    def test_missing_file(self):
        with self.assertRaises(FileNotFoundError):
            json_to_csv(str(self.json_path), str(self.csv_path))
        with self.assertRaises(FileNotFoundError):
            csv_to_xlsx(str(self.csv_path), str(self.xlsx_path))

    def test_wrong_extension(self):
        wrong = self.folder / "wrong.txt"
        wrong.write_text("name\nАнна\n", encoding="utf-8")
        with self.assertRaises(ValueError):
            csv_to_json(str(wrong), str(self.json_path))
        with self.assertRaises(ValueError):
            csv_to_xlsx(str(wrong), str(self.xlsx_path))

    def test_invalid_encoding(self):
        self.csv_path.write_bytes(b"\xff\xfe")
        with self.assertRaises(UnicodeDecodeError):
            csv_to_json(str(self.csv_path), str(self.json_path))

    def test_xlsx_content_and_width(self):
        self.csv_path.write_text("name,age\nАлександра,018\nТимур,019\n", encoding="utf-8")
        csv_to_xlsx(str(self.csv_path), str(self.xlsx_path))
        workbook = load_workbook(self.xlsx_path)
        try:
            sheet = workbook.active
            self.assertEqual(sheet.title, "Sheet1")
            self.assertEqual(list(sheet.values), [("name", "age"), ("Александра", "018"), ("Тимур", "019")])
            self.assertEqual(sheet.column_dimensions["A"].width, 12)
            self.assertEqual(sheet.column_dimensions["B"].width, 8)
        finally:
            workbook.close()

    def test_xlsx_keeps_text(self):
        self.csv_path.write_text("value\n=1+1\n#N/A\n", encoding="utf-8")
        csv_to_xlsx(str(self.csv_path), str(self.xlsx_path))
        workbook = load_workbook(self.xlsx_path)
        try:
            self.assertEqual([cell.value for cell in workbook.active["A"]], ["value", "=1+1", "#N/A"])
        finally:
            workbook.close()

    def test_csv_only_header(self):
        self.csv_path.write_text("name\n", encoding="utf-8")
        csv_to_json(str(self.csv_path), str(self.json_path))
        self.assertEqual(json.loads(self.json_path.read_text(encoding="utf-8")), [])
        csv_to_xlsx(str(self.csv_path), str(self.xlsx_path))
        workbook = load_workbook(self.xlsx_path)
        try:
            self.assertEqual(list(workbook.active.values), [("name",)])
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
